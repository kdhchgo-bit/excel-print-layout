#!/usr/bin/env python3
"""Read-only workbook structure comparison. See reported coverage limitations."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from datetime import date, datetime, time, timedelta
from pathlib import Path
import sys
import warnings
from safe_output import ensure_not_input, write_bytes

XML_FORMATS = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def value(item):
    if item is None or isinstance(item, (str, int, float, bool)):
        return item
    if isinstance(item, (date, datetime, time)):
        return {"type": type(item).__name__, "iso": item.isoformat()}
    if isinstance(item, timedelta):
        return {"type": "timedelta", "seconds": item.total_seconds()}
    if isinstance(item, bytes):
        return {"type": "bytes", "hex": item.hex()}
    # Newer openpyxl versions expose array/data-table formulas as objects.
    fields = {k: value(v) for k, v in vars(item).items() if not k.startswith("_")}
    return {"type": type(item).__name__, "fields": fields}


def xlsx_snapshot(path: Path, max_cells: int) -> dict:
    try:
        import openpyxl
        from openpyxl.utils import column_index_from_string
    except ImportError as exc:
        raise RuntimeError(".xlsx/.xlsm comparison requires openpyxl") from exc
    with warnings.catch_warnings(record=True) as found:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(path, data_only=False, keep_links=True)
        cached = openpyxl.load_workbook(path, data_only=True, keep_links=True)
    result = {"date_epoch": str(wb.epoch), "sheets": [],
              "warnings": sorted(set(str(w.message) for w in found)), "partial": False}
    try:
        for sheet in wb._sheets:
            info = {"name": sheet.title, "visibility": sheet.sheet_state,
                    "sheet_type": type(sheet).__name__}
            if not hasattr(sheet, "_cells"):
                info["unsupported"] = "chart or non-worksheet content requires native Excel review"
                result["partial"] = True
                result["sheets"].append(info)
                continue
            info["default_row_hidden"] = bool(sheet.sheet_format.zeroHeight)
            info["hidden_rows"] = sorted(int(i) for i, dim in sheet.row_dimensions.items() if dim.hidden)
            if info["default_row_hidden"]:
                info["explicit_unhidden_rows"] = sorted(int(i) for i, dim in sheet.row_dimensions.items()
                                                        if not dim.hidden)
            hidden = set()
            for key, dim in sheet.column_dimensions.items():
                first = dim.min or column_index_from_string(key)
                last = dim.max or first
                indices = set(range(first, last + 1))
                if dim.hidden:
                    hidden.update(indices)
                else:
                    hidden.difference_update(indices)
            info["hidden_cols"] = sorted(hidden)
            info["merges"] = sorted(str(item) for item in sheet.merged_cells.ranges)
            if len(sheet._cells) > max_cells:
                info["content_skipped"] = f"{len(sheet._cells)} stored cells exceeds --max-cells"
                result["partial"] = True
            else:
                cells, formula_cache = {}, {}
                cache_sheet = cached[sheet.title]
                for (row, col), cell in sheet._cells.items():
                    if cell.value is None:
                        continue
                    coordinate = f"{row},{col}"
                    cells[coordinate] = [cell.data_type, value(cell.value)]
                    if cell.data_type == "f":
                        cache = cache_sheet.cell(row, col)
                        formula_cache[coordinate] = [cache.data_type, value(cache.value)]
                info["cells"] = cells
                info["formula_cache"] = formula_cache
            result["sheets"].append(info)
    finally:
        wb.close()
        cached.close()
    return result


def xls_snapshot(path: Path, max_cells: int) -> dict:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError(".xls comparison requires xlrd with BIFF formatting support") from exc
    wb = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
    result = {"date_epoch": wb.datemode, "sheets": [], "warnings": [], "partial": False}
    try:
        for name in wb.sheet_names():
            sheet = wb.sheet_by_name(name)
            info = {"name": name, "visibility": sheet.visibility, "sheet_type": "Worksheet",
                    "default_row_hidden": bool(sheet.default_row_hidden),
                    "hidden_rows": sorted(i + 1 for i, dim in sheet.rowinfo_map.items() if dim.hidden),
                    "hidden_cols": sorted(i + 1 for i, dim in sheet.colinfo_map.items() if dim.hidden),
                    "merges": sorted([list(item) for item in sheet.merged_cells])}
            if info["default_row_hidden"]:
                info["explicit_unhidden_rows"] = sorted(i + 1 for i, dim in sheet.rowinfo_map.items()
                                                        if not dim.hidden)
            if sheet.nrows * sheet.ncols > max_cells:
                info["content_skipped"] = f"{sheet.nrows * sheet.ncols} cells exceeds --max-cells"
                result["partial"] = True
            else:
                cells = {}
                for row in range(sheet.nrows):
                    for col in range(sheet.ncols):
                        cell_type = sheet.cell_type(row, col)
                        if cell_type not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            cells[f"{row + 1},{col + 1}"] = [cell_type, value(sheet.cell_value(row, col))]
                info["cells"] = cells
            result["sheets"].append(info)
            wb.unload_sheet(name)
    finally:
        wb.release_resources()
    return result


def diff_mapping(old: dict, new: dict) -> dict:
    left, right = set(old), set(new)
    changed = sorted(key for key in left & right if old[key] != new[key])
    return {"original_only_count": len(left - right), "revised_only_count": len(right - left),
            "different_count": len(changed), "original_only_sample": sorted(left - right)[:20],
            "revised_only_sample": sorted(right - left)[:20], "different_sample": changed[:20]}


def compare(old: dict, new: dict) -> dict:
    differences = {}
    if old["date_epoch"] != new["date_epoch"]:
        differences["date_epoch"] = {"original": old["date_epoch"], "revised": new["date_epoch"]}
    a_names = [s["name"] for s in old["sheets"]]
    b_names = [s["name"] for s in new["sheets"]]
    if a_names != b_names:
        differences["sheet_order"] = {"original": a_names, "revised": b_names}
    b_sheets = {s["name"]: s for s in new["sheets"]}
    per_sheet = []
    for left in old["sheets"]:
        right = b_sheets.get(left["name"])
        entry = {"name": left["name"], "differences": {}}
        if right is None:
            entry["differences"]["missing"] = True
        else:
            for key in ("visibility", "sheet_type", "default_row_hidden", "hidden_rows", "hidden_cols",
                        "explicit_unhidden_rows", "merges"):
                if left.get(key) != right.get(key):
                    entry["differences"][key] = {"original": left.get(key), "revised": right.get(key)}
            for key in ("cells", "formula_cache"):
                if key in left and key in right and left[key] != right[key]:
                    entry["differences"][key] = diff_mapping(left[key], right[key])
        entry["status"] = "different" if entry["differences"] else "match_within_checked_fields"
        per_sheet.append(entry)
    partial = old["partial"] or new["partial"]
    different = bool(differences) or any(item["differences"] for item in per_sheet)
    return {"status": "different" if different else ("incomplete" if partial else "match_within_checked_fields"),
            "differences": differences, "sheets": per_sheet,
            "warnings": {"original": old["warnings"], "revised": new["warnings"]},
            "partial": partial}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("original", type=Path)
    parser.add_argument("revised", type=Path)
    parser.add_argument("--output", type=Path, help="write JSON report instead of stdout")
    parser.add_argument("--overwrite", action="store_true", help="replace a different existing report")
    parser.add_argument("--max-cells", type=int, default=2_000_000,
                        help="max stored cells per sheet; larger sheets are reported incomplete")
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        ext = args.original.suffix.lower()
        if args.max_cells < 1 or ext != args.revised.suffix.lower():
            raise ValueError("use matching workbook extensions and a positive --max-cells")
        if os.path.samefile(args.original, args.revised):
            raise ValueError("original and revised must be different files")
        if args.output:
            ensure_not_input(args.output, [args.original, args.revised])
        original_hash, revised_hash = sha(args.original), sha(args.revised)
        if ext == ".xls":
            old, new = xls_snapshot(args.original, args.max_cells), xls_snapshot(args.revised, args.max_cells)
            limitations = ["BIFF formula expressions and external-link targets are not checked by xlrd",
                           "non-worksheet sheets and drawings require native Excel inspection",
                           "print layout, cell clipping and images are not visually checked"]
        elif ext in XML_FORMATS:
            old, new = xlsx_snapshot(args.original, args.max_cells), xlsx_snapshot(args.revised, args.max_cells)
            limitations = ["external-link targets require native Excel or dedicated link inspection",
                           "drawings, print layout, cell clipping and images are not visually checked",
                           "macro/VBA contents are not compared"]
        else:
            raise ValueError("supported: .xls, .xlsx, .xlsm, .xltx, .xltm")
        report = compare(old, new)
        if sha(args.original) != original_hash or sha(args.revised) != revised_hash:
            report["status"] = "incomplete"
            report["partial"] = True
            report["race"] = "input file changed during comparison"
        report.update({"original_sha256": original_hash, "revised_sha256": revised_hash,
                       "format": ext, "limitations": limitations})
        text = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        if args.output:
            write_bytes(args.output, text.encode("utf-8"), overwrite=args.overwrite)
        else:
            print(text, end="")
        return {"match_within_checked_fields": 0, "different": 1, "incomplete": 2}[report["status"]]
    except Exception as exc:
        print(json.dumps({"status": "error", "error": f"{type(exc).__name__}: {exc}"},
                         ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
