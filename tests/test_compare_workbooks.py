import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "compare_workbooks.py"
HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
HAS_XL = importlib.util.find_spec("xlrd") is not None and importlib.util.find_spec("xlwt") is not None


class StructuralComparisonTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def compare(self, original, revised):
        report = self.root / f"report-{original.stem}-{revised.stem}.json"
        result = subprocess.run([sys.executable, str(SCRIPT), str(original), str(revised), "--output", str(report)],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertTrue(report.exists(), result.stdout + result.stderr)
        return result.returncode, json.loads(report.read_text(encoding="utf-8"))

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl not installed")
    def test_ooxml_expands_hidden_column_ranges(self):
        import openpyxl
        source = self.root / "original.xlsx"
        copy = self.root / "copy.xlsx"
        changed = self.root / "changed.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Visible"
        ws["A1"] = "header"
        ws["A2"] = "=1+2"
        ws.merge_cells("E1:F1")
        ws.row_dimensions[3].hidden = True
        ws.column_dimensions.group("B", "D", hidden=True)
        wb.create_sheet("Hidden").sheet_state = "hidden"
        wb.save(source)
        shutil.copyfile(source, copy)
        code, report = self.compare(source, copy)
        self.assertEqual((code, report["status"]), (0, "match_within_checked_fields"))
        original_bytes = source.read_bytes()
        blocked = subprocess.run([sys.executable, str(SCRIPT), str(source), str(copy), "--output", str(source)],
                                 capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(blocked.returncode, 2)
        self.assertEqual(source.read_bytes(), original_bytes)
        unrelated = self.root / "existing.json"
        unrelated.write_text("do not replace", encoding="utf-8")
        blocked = subprocess.run([sys.executable, str(SCRIPT), str(source), str(copy), "--output", str(unrelated)],
                                 capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(blocked.returncode, 2)
        self.assertEqual(unrelated.read_text(encoding="utf-8"), "do not replace")
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        with zipfile.ZipFile(source) as src, zipfile.ZipFile(changed, "w") as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    cols = root.find(ns + "cols")
                    first = cols.find(ns + "col")
                    self.assertEqual((first.attrib["min"], first.attrib["max"]), ("2", "4"))
                    first.set("max", "2")
                    ET.SubElement(cols, ns + "col", {"min": "3", "max": "4", "hidden": "0"})
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                dst.writestr(item, data)
        code, report = self.compare(source, changed)
        self.assertEqual(code, 1)
        delta = report["sheets"][0]["differences"]["hidden_cols"]
        self.assertEqual(delta["original"], [2, 3, 4])
        self.assertEqual(delta["revised"], [2])

        # A forced low limit must not turn an uninspected sheet into a pass.
        report = self.root / "limited.json"
        result = subprocess.run([sys.executable, str(SCRIPT), str(source), str(copy), "--max-cells", "1",
                                 "--output", str(report)], capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 2)
        self.assertEqual(json.loads(report.read_text(encoding="utf-8"))["status"], "incomplete")

    @unittest.skipUnless(HAS_XL, "xlrd/xlwt not installed")
    def test_legacy_xls_detects_hidden_and_value_changes(self):
        import xlwt
        source, changed = self.root / "a.xls", self.root / "b.xls"
        for path, hidden, amount in ((source, True, 100), (changed, False, 101)):
            wb = xlwt.Workbook()
            ws = wb.add_sheet("Sheet1")
            ws.write(0, 0, amount)
            ws.row(3).hidden = hidden
            ws.col(2).hidden = True
            wb.save(str(path))
        code, report = self.compare(source, changed)
        self.assertEqual(code, 1)
        self.assertIn("hidden_rows", report["sheets"][0]["differences"])
        self.assertEqual(report["sheets"][0]["differences"]["cells"]["different_count"], 1)
        self.assertIn("BIFF formula", report["limitations"][0])


if __name__ == "__main__":
    unittest.main()
